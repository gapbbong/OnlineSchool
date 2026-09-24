from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from tests.conftest import auth_headers


@pytest.mark.asyncio
async def test_bulk_import_template_reflects_school_departments(client):
    headers = await auth_headers("hong@kse.hs.kr")
    res = await client.get("/api/v1/teachers/onboarding/template", headers=headers)
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]

    wb = load_workbook(BytesIO(res.content))
    assert "교무부" in [row[0] for row in wb["참고-부서목록"].iter_rows(min_row=2, values_only=True)]


@pytest.mark.asyncio
async def test_bulk_import_partial_success(client):
    headers = await auth_headers("hong@kse.hs.kr")

    wb = Workbook()
    ws = wb.active
    ws.title = "교사일괄등록"
    ws.append([
        "이름*", "휴대폰번호*", "Workspace 이메일*", "차량번호", "부서명*",
        "직책", "담당업무", "담임학년", "담임반", "담당과목명",
        "권한(TEACHER/STAFF/DEPARTMENT_HEAD)",
    ])
    # 정상 행
    ws.append(["엑셀교사1", "010-1111-2222", "excel1@kse.hs.kr", "", "교무부", "교과교사", "", "", "", "", "TEACHER"])
    # 존재하지 않는 부서 -> 실패해야 함
    ws.append(["엑셀교사2", "010-3333-4444", "excel2@kse.hs.kr", "", "없는부서", "교과교사", "", "", "", "", "TEACHER"])
    # 필수값 누락 -> 실패해야 함
    ws.append(["", "", "", "", "", "", "", "", "", "", "TEACHER"])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    res = await client.post(
        "/api/v1/teachers/onboarding/bulk",
        headers=headers,
        files={"file": ("bulk.xlsx", buffer.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert res.status_code == 200
    data = res.json()
    assert data["total_rows"] == 3
    assert data["succeeded"] == 1
    assert data["failed"] == 2
    assert any(r["success"] and r["name"] == "엑셀교사1" for r in data["results"])
    assert any(not r["success"] and "부서" in (r["error"] or "") for r in data["results"])
    assert any(not r["success"] and "누락" in (r["error"] or "") for r in data["results"])
